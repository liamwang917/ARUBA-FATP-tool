"""Locked V13.6 workbook generation."""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .config import METADATA_CONTEXT_COLUMNS, SHEET_NAMES
from .models import Curve, TestRun


FREQUENCY_LEADS = {
    "02_FR_original": ("SN", "Test_Time", "RawData_Match_Status"),
    "03_FR_1_3": ("SN", "Test_Time", "Path_Result", "FR_File_Result"),
    "04_FR_1_12": ("SN", "Test_Time", "RawData_Match_Status"),
    "05_THD": ("SN", "Test_Time", "Path_Result", "Result"),
    "06_Phase": ("SN", "Test_Time", "Path_Result", "Result"),
    "07_Noise": ("SN", "Test_Time", "Path_Result", "Result"),
}


def _key(value: object) -> str:
    return str(value or "").strip().casefold()


def _metadata_item_columns(runs: list[TestRun]) -> list[tuple[str, str]]:
    columns: list[tuple[str, str]] = []
    seen = set()
    for run in runs:
        for display_name, _ in run.station_items:
            normalized = _key(display_name)
            if normalized and normalized not in seen:
                seen.add(normalized)
                columns.append((normalized, display_name))
    return columns


def _metadata_row(run: TestRun, item_columns: list[tuple[str, str]]) -> list[object]:
    context = {
        "Test_Type": run.test_type,
        "Mode": run.mode,
        "SN": run.sn,
        "Run_ID": run.run_id,
        "Test_Time": run.main_timestamp,
        "Station": run.station,
        "Path_Result": run.path_result,
        "FR_File_Result": run.fr_file_result,
        "Latest_Run": "TRUE" if run.latest_run else "FALSE",
        "RawData_Result": run.rawdata_result,
        "RawData_Match_Status": run.rawdata_match_status,
        "RawData_Time_Delta_s": run.rawdata_time_delta_s,
        "RawData_PF_Mismatch": (
            run.rawdata_pf_mismatch if run.rawdata_pf_mismatch != "" else ""
        ),
        "Main_CSV": run.main_csv,
        "FR_CSV": run.fr_csv,
        "Noise_CSV": run.noise_csv,
    }
    station_values = {_key(name): value for name, value in run.station_items}
    return (
        [context[column] for column in METADATA_CONTEXT_COLUMNS]
        + [station_values.get(normalized, "") for normalized, _ in item_columns]
    )


def _curve_and_lead(run: TestRun, sheet: str) -> tuple[Curve, list[object]]:
    if sheet == "02_FR_original":
        return run.raw_fr_original, [
            run.sn, run.rawdata_timestamp, run.rawdata_match_status
        ]
    if sheet == "03_FR_1_3":
        return run.fr, [
            run.sn, run.fr_timestamp, run.path_result, run.fr_file_result
        ]
    if sheet == "04_FR_1_12":
        return run.raw_fr_1_12, [
            run.sn, run.rawdata_timestamp, run.rawdata_match_status
        ]
    mapping = {
        "05_THD": (run.thd, "THD", run.fr_timestamp),
        "06_Phase": (run.phase, "PHASE", run.fr_timestamp),
        "07_Noise": (run.noise, "NOISE", run.noise_timestamp),
    }
    curve, result_key, timestamp = mapping[sheet]
    return curve, [
        run.sn, timestamp, run.path_result, run.item_results.get(result_key, "")
    ]


def _write_curve_sheet(
    ws, runs: list[TestRun], sheet: str, logger: logging.Logger
) -> None:
    lead = FREQUENCY_LEADS[sheet]
    axis = next(
        (
            curve.frequencies
            for run in runs
            for curve, _ in [_curve_and_lead(run, sheet)]
            if curve.frequencies
        ),
        [],
    )
    ws.append(list(lead) + axis)
    for run in runs:
        curve, leading = _curve_and_lead(run, sheet)
        values: list[object] = []
        if curve.frequencies:
            if curve.frequencies == axis:
                values = curve.values
            else:
                logger.warning(
                    "Frequency-axis mismatch on %s for %s; values left blank",
                    sheet,
                    run.run_path,
                )
        ws.append(leading + values)


def _write_scalar_sheets(wb: Workbook, runs: list[TestRun]) -> None:
    snr = wb["08_SNR"]
    snr.append([
        "SN", "Test_Time", "Path_Result", "Result", "SNR_dB",
        "SNR_Source_Status",
    ])
    sensitivity = wb["09_Sensitivity"]
    sensitivity.append([
        "SN", "Test_Time", "Path_Result", "Result", "Sensitivity_dBFS",
        "Sensitivity_Source_Status",
    ])
    for run in runs:
        snr_status = (
            "Present in Main Station CSV"
            if run.snr_present
            else "Missing in Main Station CSV"
        )
        sensitivity_status = (
            "Present in Main Station CSV"
            if run.sensitivity_present
            else "Missing in Main Station CSV"
        )
        snr.append([
            run.sn,
            run.main_timestamp,
            run.path_result,
            run.item_results.get("SNR", ""),
            run.snr,
            snr_status,
        ])
        sensitivity.append([
            run.sn,
            run.main_timestamp,
            run.path_result,
            run.item_results.get("SENSITIVITY", ""),
            run.sensitivity,
            sensitivity_status,
        ])


def _format_workbook(wb: Workbook) -> None:
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, datetime):
                    cell.number_format = (
                        "yyyy-mm-dd hh:mm:ss.000"
                        if cell.value.microsecond
                        else "yyyy-mm-dd hh:mm:ss"
                    )
        for column in ws.columns:
            letter = column[0].column_letter
            width = max(len(str(cell.value or "")) for cell in column) + 2
            ws.column_dimensions[letter].width = min(28, max(10, width))


def build_workbook(
    runs: list[TestRun], output_path: Path, logger: logging.Logger
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEET_NAMES:
        wb.create_sheet(name)

    item_columns = _metadata_item_columns(runs)
    metadata = wb["01_Metadata"]
    metadata.append(
        list(METADATA_CONTEXT_COLUMNS)
        + [display_name for _, display_name in item_columns]
    )
    for run in runs:
        metadata.append(_metadata_row(run, item_columns))
    for sheet in SHEET_NAMES[1:7]:
        _write_curve_sheet(wb[sheet], runs, sheet, logger)
    _write_scalar_sheets(wb, runs)
    _format_workbook(wb)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

