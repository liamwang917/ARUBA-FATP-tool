"""Surgical V14.5 report-template population; never round-trips the master."""

from __future__ import annotations

import hashlib
import io
import logging
import os
import re
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


TEMPLATE_SHA256 = "a556aa066b9412fb17512bb4f992d26909676c5bc53a138616e2cebef2925724"
MAX_DUTS = 947
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
ET.register_namespace("", NS)
ET.register_namespace("r", REL_NS)

SHEETS = {
    "03_FR_1_3": ("Frequency Response_1_3", "FR_File_Result", 986),
    "04_FR_1_12": ("Frequency Response_1_12", "RawData_Match_Status", 986),
    "02_FR_original": ("Frequency Response_orignal", "RawData_Match_Status", 986),
    "05_THD": ("THD", "Result", 986),
    "06_Phase": ("Phase", "Result", 986),
    "07_Noise": ("Noise Floor", "Result", 987),
}
SCALARS = {"08_SNR": ("SNR", "SNR_dB", 989), "09_Sensitivity": ("Sensitivity", "Sensitivity_dBFS", 989)}
CURVE_CHART_SHEETS = (
    "Frequency Response_1_3", "Frequency Response_1_12",
    "Frequency Response_orignal", "THD", "Phase", "Noise Floor",
)


class V14ReportError(ValueError):
    pass


def _xml(payload: bytes) -> ET.Element:
    """Parse after registering every namespace used by the source OOXML part.

    Excel's mc:Ignorable values refer to literal prefixes.  Registering the
    original bindings before serialization keeps those prefix references valid.
    """
    for _, (prefix, uri) in ET.iterparse(io.BytesIO(payload), events=("start-ns",)):
        if prefix not in {"xml", "xmlns"} and not re.fullmatch(r"ns\d+", prefix):
            ET.register_namespace(prefix, uri)
    return ET.fromstring(payload)


def _serialize(root: ET.Element, original: bytes | None = None) -> bytes:
    """Serialize while retaining root-level namespace declarations verbatim."""
    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    if not original:
        return payload
    original_tag = re.search(rb"<[^!?][^>]*>", original).group(0)
    declarations = re.findall(rb"\s(xmlns(?::[^=\s]+)?=\"[^\"]+\")", original_tag)
    generated_tag = re.search(rb"<[^!?][^>]*>", payload).group(0)
    additions = [item for item in declarations if item not in generated_tag]
    if additions:
        payload = payload.replace(generated_tag, generated_tag[:-1] + b" " + b" ".join(additions) + b">", 1)
    return payload


def verify_template(template: Path, expected_sha256: str = TEMPLATE_SHA256) -> None:
    if not template.is_file():
        raise V14ReportError(
            "Approved V14.5 template is missing. Place "
            "Post-MIC limit_20260911.xlsx at templates/"
        )
    digest = hashlib.sha256(template.read_bytes()).hexdigest()
    if digest != expected_sha256:
        raise V14ReportError(
            "Approved V14.5 template hash is incorrect. Replace templates/"
            "Post-MIC limit_20260911.xlsx with the approved master"
        )


def _column(index: int) -> str:
    text = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        text = chr(65 + remainder) + text
    return text


def _xml_column(reference: str) -> int:
    result = 0
    for character in re.match(r"[A-Z]+", reference).group(0):
        result = result * 26 + ord(character) - 64
    return result


def _row_cells(row):
    values = {}
    for cell in row.findall(f"{{{NS}}}c"):
        value = cell.findtext(f"{{{NS}}}v")
        if value is not None:
            try:
                value = float(value)
                value = int(value) if value.is_integer() else value
            except ValueError:
                pass
            values[_xml_column(cell.attrib["r"])] = value
    return values


def _read_summary(path: Path) -> dict[str, tuple[list[object], list[dict[str, object]]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        result = {}
        for ws in wb.worksheets:
            iterator = ws.iter_rows(values_only=True)
            header = list(next(iterator, ()))
            result[ws.title] = (header, [dict(zip(header, values)) for values in iterator])
        return result
    finally:
        wb.close()


def _normal(value: object) -> str:
    return str(value if value is not None else "").strip().casefold()


def _display(value: object) -> object:
    text = str(value).strip() if value is not None else ""
    return text.upper() if text.casefold() in {"pass", "fail"} else value


def _cell(row: ET.Element, column: int, value: object, *, compact_datetime: bool = True) -> None:
    reference = f"{_column(column)}{row.attrib['r']}"
    cell = next((item for item in row.findall(f"{{{NS}}}c") if item.attrib.get("r") == reference), None)
    if cell is None:
        cell = ET.SubElement(row, f"{{{NS}}}c", {"r": reference})
    for child in list(cell):
        cell.remove(child)
    if value is None or value == "":
        cell.attrib.pop("t", None)
        return
    if isinstance(value, datetime):
        value = (value.strftime("%Y%m%d%H%M%S") if compact_datetime else
                 value.strftime("%Y-%m-%d %H:%M:%S") + f".{value.microsecond // 1000:03d}")
    if isinstance(value, bool):
        value = int(value)
    if isinstance(value, (int, float)):
        cell.attrib.pop("t", None)
        ET.SubElement(cell, f"{{{NS}}}v").text = str(value)
    else:
        cell.attrib["t"] = "inlineStr"
        inline = ET.SubElement(cell, f"{{{NS}}}is")
        ET.SubElement(inline, f"{{{NS}}}t").text = str(_display(value))


def _sheet_path(package: zipfile.ZipFile) -> tuple[ET.Element, ET.Element, dict[str, str]]:
    workbook = _xml(package.read("xl/workbook.xml"))
    rels = _xml(package.read("xl/_rels/workbook.xml.rels"))
    targets = {rel.attrib["Id"]: rel.attrib["Target"].lstrip("/") for rel in rels}
    paths = {}
    for sheet in workbook.findall(f".//{{{NS}}}sheet"):
        target = targets[sheet.attrib[f"{{{REL_NS}}}id"]]
        paths[sheet.attrib["name"]] = target if target.startswith("xl/") else f"xl/{target}"
    return workbook, rels, paths


def _ensure_row(root: ET.Element, number: int) -> ET.Element:
    data = root.find(f"{{{NS}}}sheetData")
    row = next((item for item in data.findall(f"{{{NS}}}row") if item.attrib.get("r") == str(number)), None)
    if row is None:
        row = ET.SubElement(data, f"{{{NS}}}row", {"r": str(number)})
        for index, existing in enumerate(list(data)):
            if existing is row:
                break
            if int(existing.attrib.get("r", "0")) > number:
                data.remove(row)
                data.insert(index, row)
                break
    return row


def _clear_existing(root: ET.Element, first_column: int, last_column: int, end_row: int) -> None:
    """Clear legacy values only; never manufacture a full-capacity blank grid."""
    data = root.find(f"{{{NS}}}sheetData")
    for row in data.findall(f"{{{NS}}}row"):
        if not 40 <= int(row.attrib.get("r", "0")) <= end_row:
            continue
        for cell in row.findall(f"{{{NS}}}c"):
            column = _xml_column(cell.attrib["r"])
            if first_column <= column <= last_column:
                for child in list(cell):
                    cell.remove(child)
                cell.attrib.pop("t", None)


def _set_dimension(root: ET.Element, required_last_column: int, required_last_row: int) -> None:
    dimension = root.find(f"{{{NS}}}dimension")
    ref = dimension.attrib.get("ref", "A1")
    match = re.fullmatch(r"([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?", ref)
    if not match:
        return
    start_column, start_row, end_column, end_row = match.groups()
    end_column = end_column or start_column
    end_row = int(end_row or start_row)
    end_index = max(_xml_column(end_column + "1"), required_last_column)
    dimension.attrib["ref"] = f"{start_column}{start_row}:{_column(end_index)}{max(end_row, required_last_row)}"


def _clear_and_write_curve(xml: bytes, header: list[object], rows: list[dict[str, object]], context: str, end_row: int) -> bytes:
    root = _xml(xml)
    data = root.find(f"{{{NS}}}sheetData")
    template_axis = _row_cells(next(row for row in data.findall(f"{{{NS}}}row") if row.attrib.get("r") == "39"))
    source_axis = header[header.index(context) + 1:]
    destination_axis = [template_axis[column] for column in sorted(template_axis) if column >= 4]
    is_rawdata = context == "RawData_Match_Status"
    no_rawdata_curve = not source_axis and is_rawdata and all(
        row.get(context) in {"NOT_PROVIDED", "UNMATCHED", "AMBIGUOUS", ""} for row in rows
    )
    if not no_rawdata_curve and [_normal(x) for x in source_axis] != [_normal(x) for x in destination_axis]:
        raise V14ReportError(f"Frequency-axis mismatch for {context}")
    _clear_existing(root, 1, len(destination_axis) + 3, end_row)
    for offset, source in enumerate(rows):
        row = _ensure_row(root, 40 + offset)
        _cell(row, 1, source.get("SN"))
        _cell(row, 2, source.get("Test_Time"))
        _cell(row, 3, source.get(context))
        for index, frequency in enumerate(source_axis, 4):
            _cell(row, index, source.get(frequency))
    _set_dimension(root, len(destination_axis) + 3, 39 + len(rows))
    return _serialize(root, xml)


def _clear_and_write_scalar(xml: bytes, rows: list[dict[str, object]], value_key: str, end_row: int) -> bytes:
    root = _xml(xml)
    _clear_existing(root, 1, 2, end_row)
    for offset, source in enumerate(rows):
        row = _ensure_row(root, 40 + offset)
        _cell(row, 1, source.get("SN")); _cell(row, 2, source.get(value_key))
    _set_dimension(root, 2, 39 + len(rows))
    return _serialize(root, xml)


def _metadata_xml(header: list[object], rows: list[dict[str, object]]) -> bytes:
    root = ET.Element(f"{{{NS}}}worksheet")
    ET.SubElement(root, f"{{{NS}}}dimension", {"ref": f"A1:{_column(max(1, len(header)))}{len(rows) + 1}"})
    data = ET.SubElement(root, f"{{{NS}}}sheetData")
    for number, values in enumerate([dict(zip(header, header))] + rows, 1):
        row = ET.SubElement(data, f"{{{NS}}}row", {"r": str(number)})
        for column, key in enumerate(header, 1):
            _cell(row, column, values.get(key), compact_datetime=False)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _add_metadata(workbook, rels, content_types, header, rows, workbook_source: bytes) -> tuple[bytes, bytes, bytes, bytes]:
    sheets = workbook.find(f"{{{NS}}}sheets")
    ids = [int(sheet.attrib["sheetId"]) for sheet in sheets]
    existing = [int(re.search(r"rId(\d+)", rel.attrib["Id"]).group(1)) for rel in rels if re.fullmatch(r"rId\d+", rel.attrib["Id"])]
    rid = f"rId{max(existing, default=0) + 1}"
    sheet = ET.SubElement(sheets, f"{{{NS}}}sheet", {"name": "Metadata", "sheetId": str(max(ids) + 1), f"{{{REL_NS}}}id": rid})
    ET.SubElement(rels, f"{{{PKG_REL_NS}}}Relationship", {"Id": rid, "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet", "Target": "worksheets/sheet12.xml"})
    ET.SubElement(content_types, f"{{{CT_NS}}}Override", {"PartName": "/xl/worksheets/sheet12.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"})
    return (_serialize(workbook, workbook_source), ET.tostring(rels, encoding="utf-8", xml_declaration=True), ET.tostring(content_types, encoding="utf-8", xml_declaration=True), _metadata_xml(header, rows))


def _apply_v145_limit_correction(xml: bytes) -> bytes:
    root = _xml(xml)
    data = root.find(f"{{{NS}}}sheetData")
    for row_number, formula in ((35, "D28+D32"), (36, "D28-D32")):
        row = _ensure_row(root, row_number)
        cell = next((item for item in row.findall(f"{{{NS}}}c") if item.attrib.get("r") == f"D{row_number}"), None)
        if cell is None:
            cell = ET.SubElement(row, f"{{{NS}}}c", {"r": f"D{row_number}"})
        for child in list(cell):
            cell.remove(child)
        cell.attrib.pop("t", None)
        ET.SubElement(cell, f"{{{NS}}}f").text = formula
    return _serialize(root, xml)


def _apply_chart_source_updates(payload: bytes, sheet: str) -> bytes:
    """Apply the two explicitly approved V14.5 chart-source exceptions only."""
    text = payload.decode("utf-8")
    escaped = re.escape(sheet)
    sheet_reference = rf"(?:'{escaped}'|{escaped})!"

    def sn_only(match: re.Match) -> str:
        row = int(match.group("row"))
        return match.group(0) if row < 40 else f"{match.group('sheet_ref')}$A${row}"

    text = re.sub(
        rf"(?P<sheet_ref>{sheet_reference})\$A\$(?P<row>\d+):\$C\$(?P=row)", sn_only, text,
    )
    if sheet == "Noise Floor":
        text = re.sub(
            r"'Noise Floor'!\$D\$(?P<row>\d+):\$ADW\$(?P=row)",
            lambda match: (match.group(0) if int(match.group("row")) < 39
                           else f"'Noise Floor'!$N${match.group('row')}:$ADW${match.group('row')}"),
            text,
        )
    return text.encode("utf-8")


def _chart_capacity(package: zipfile.ZipFile) -> int:
    capacities = []
    for name in package.namelist():
        if name.startswith("xl/charts/") and name.endswith(".xml"):
            numbers = [int(x) for x in re.findall(r"\$(?:[A-Z]+)\$(\d+)", package.read(name).decode("utf-8", "ignore"))]
            coverage = max((number - 39 for number in numbers if number >= 40), default=None)
            if coverage is not None:
                capacities.append(coverage)
    return min(capacities, default=MAX_DUTS)


def build_v14_report(template: Path, summaries: list[Path], output: Path, logger: logging.Logger | None = None, *, verify_hash: bool = True) -> Path:
    """Populate a template copy using V13.6 normalized summary workbook(s)."""
    logger = logger or logging.getLogger("aruba_fatp_v14")
    if verify_hash:
        verify_template(template)
    source = {}
    for summary in summaries:
        for name, table in _read_summary(summary).items():
            if name not in source:
                source[name] = (table[0], list(table[1]))
            else:
                if source[name][0] != table[0]:
                    raise V14ReportError(f"Summary header mismatch for {name}")
                source[name][1].extend(table[1])
    count = len(source["01_Metadata"][1])
    if count > MAX_DUTS:
        raise V14ReportError(f"V14 report has {count} DUTs; maximum is {MAX_DUTS}")
    output.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(prefix=f".{output.stem}.", suffix=".xlsx", dir=output.parent)
    os.close(handle)
    Path(temporary_name).unlink(missing_ok=True)
    temporary = Path(temporary_name)
    try:
      shutil.copyfile(template, temporary)
      with zipfile.ZipFile(template) as original:
        if any(name.startswith("xl/externalLinks/") for name in original.namelist()):
            raise V14ReportError("Approved master must not contain external workbook links")
        workbook_source = original.read("xl/workbook.xml")
        workbook, rels, paths = _sheet_path(original)
        content_types = _xml(original.read("[Content_Types].xml"))
        changes = {}
        for source_sheet, (target, context, end_row) in SHEETS.items():
            header, rows = source[source_sheet]
            changes[paths[target]] = _clear_and_write_curve(original.read(paths[target]), header, rows, context, end_row)
        changes[paths["Frequency Response_1_3"]] = _apply_v145_limit_correction(changes[paths["Frequency Response_1_3"]])
        for source_sheet, (target, value, end_row) in SCALARS.items():
            _, rows = source[source_sheet]
            changes[paths[target]] = _clear_and_write_scalar(original.read(paths[target]), rows, value, end_row)
        for sheet in workbook.findall(f".//{{{NS}}}sheet"):
            if sheet.attrib["name"] in {"Frequency Response_1_12", "Frequency Response_orignal"}:
                sheet.attrib.pop("state", None)
        calc = workbook.find(f"{{{NS}}}calcPr")
        if calc is None:
            calc = ET.SubElement(workbook, f"{{{NS}}}calcPr")
        calc.attrib.update({"calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1"})
        header, rows = source["01_Metadata"]
        changes["xl/workbook.xml"], changes["xl/_rels/workbook.xml.rels"], changes["[Content_Types].xml"], changes["xl/worksheets/sheet12.xml"] = _add_metadata(workbook, rels, content_types, header, rows, workbook_source)
        for chart_name in (name for name in original.namelist() if name.startswith("xl/charts/") and name.endswith(".xml")):
            chart = original.read(chart_name)
            for sheet in CURVE_CHART_SHEETS:
                if re.search(rf"(?:'{re.escape(sheet)}'|{re.escape(sheet)})!", chart.decode("utf-8", "ignore")):
                    changes[chart_name] = _apply_chart_source_updates(chart, sheet)
                    break
        if count > _chart_capacity(original):
            logger.warning("Chart display coverage is smaller than %d DUTs; charts are unchanged", count)
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as destination:
            for item in original.infolist():
                if item.filename in {"xl/calcChain.xml"}:
                    continue
                payload = changes.get(item.filename, original.read(item.filename))
                if item.filename == "xl/_rels/workbook.xml.rels":
                    rel_root = _xml(payload)
                    for rel in list(rel_root):
                        if rel.attrib.get("Type", "").endswith("/calcChain"):
                            rel_root.remove(rel)
                    payload = ET.tostring(rel_root, encoding="utf-8", xml_declaration=True)
                if item.filename == "[Content_Types].xml":
                    types = _xml(payload)
                    for override in list(types):
                        if override.attrib.get("PartName") == "/xl/calcChain.xml":
                            types.remove(override)
                    payload = ET.tostring(types, encoding="utf-8", xml_declaration=True)
                destination.writestr(item, payload)
            destination.writestr("xl/worksheets/sheet12.xml", changes["xl/worksheets/sheet12.xml"])
      temporary.replace(output)
      return output
    finally:
      temporary.unlink(missing_ok=True)
