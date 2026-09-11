import logging
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from src.config import SHEET_NAMES
from src.v14_report import MAX_DUTS, V14ReportError, build_v14_report
from src import v14_main


TARGETS = [
    "Audio Limit for EV", "Audio Limit for EV3", "Audio Limit_Sigma",
    "Frequency Response_1_3", "Frequency Response_1_12",
    "Frequency Response_orignal", "THD", "Phase", "Noise Floor", "SNR", "Sensitivity",
]
CURVES = {
    "03_FR_1_3": ("Frequency Response_1_3", "FR_File_Result"),
    "04_FR_1_12": ("Frequency Response_1_12", "RawData_Match_Status"),
    "02_FR_original": ("Frequency Response_orignal", "RawData_Match_Status"),
    "05_THD": ("THD", "Result"), "06_Phase": ("Phase", "Result"), "07_Noise": ("Noise Floor", "Result"),
}


def make_summary(path, rows=1, axis=(100, 200), mismatch=False, rawdata_missing=False):
    wb = Workbook(); wb.remove(wb.active)
    for name in SHEET_NAMES: wb.create_sheet(name)
    metadata = wb["01_Metadata"]
    metadata.append(["SN", "Test_Time", "synthetic_key"])
    for i in range(rows): metadata.append([f"SYNTH-{i:04d}", "2026-09-11", f"x{i}"])
    for source, (_, context) in CURVES.items():
        ws = wb[source]
        raw_source = source in {"02_FR_original", "04_FR_1_12"}
        ws.append(["SN", "Test_Time", context] if rawdata_missing and raw_source else ["SN", "Test_Time", context, *axis])
        for i in range(rows):
            values = [f"SYNTH-{i:04d}", "2026-09-11", "NOT_PROVIDED" if raw_source else "pass"]
            ws.append(values if rawdata_missing and raw_source else values + [-1.0, -2.0])
    for source, value in (("08_SNR", "SNR_dB"), ("09_Sensitivity", "Sensitivity_dBFS")):
        ws = wb[source]; ws.append(["SN", "Test_Time", "Path_Result", "Result", value])
        for i in range(rows): ws.append([f"SYNTH-{i:04d}", "2026-09-11", "PASS", "FAIL", 31.5 if source == "08_SNR" else -31.5])
    if mismatch: wb["03_FR_1_3"].cell(1, 5).value = 300
    wb.save(path)


def make_template(path, axis=(100, 200), chart_capacity=1):
    wb = Workbook(); first = wb.active; first.title = TARGETS[0]
    for name in TARGETS[1:]: wb.create_sheet(name)
    for name in TARGETS[3:9]:
        ws = wb[name]; ws.sheet_state = "hidden" if name.startswith("Frequency Response_") and name != "Frequency Response_1_3" else "visible"
        ws.cell(39, 1).value = "SN"; ws.cell(39, 2).value = "Test_Time"; ws.cell(39, 3).value = "Result"
        for col, value in enumerate(axis, 4): ws.cell(39, col).value = value
        ws.cell(40, 1).value = "LEGACY"; ws.cell(40, 4).value = 999
    for name in ("SNR", "Sensitivity"):
        ws = wb[name]; ws.cell(39, 1).value = "SN"; ws.cell(39, 2).value = "Value"; ws.cell(40, 1).value = "LEGACY"; ws.cell(40, 2).value = "=1+1"
    wb.save(path)
    # Add representative protected parts, chart capacity, and calcChain. The generator must preserve/remove them correctly.
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as package:
        package.writestr("xl/charts/chart1.xml", f"<chart><f>Sheet!$A$40:$A${39 + chart_capacity}</f></chart>")
        package.writestr("xl/drawings/drawing1.xml", b"<drawing synthetic='true'/>")
        package.writestr("xl/theme/theme1.xml", b"<theme synthetic='true'/>")
        package.writestr("xl/printerSettings/printerSettings1.bin", b"synthetic-printer")
        package.writestr("xl/calcChain.xml", b"<calcChain/>")
        rels = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
        ET.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {"Id":"rId99","Type":"http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain","Target":"calcChain.xml"})
        package.writestr("xl/_rels/workbook.xml.rels", ET.tostring(rels, encoding="utf-8", xml_declaration=True))
        types = ET.fromstring(package.read("[Content_Types].xml"))
        ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {"PartName":"/xl/calcChain.xml","ContentType":"application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"})
        package.writestr("[Content_Types].xml", ET.tostring(types, encoding="utf-8", xml_declaration=True))


def add_excel_compatibility_namespaces(path):
    """Add real-master-like compatibility prefixes without using the V14 writer."""
    namespaces = ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main" xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision"'
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as package:
        for name, ignorable, extension in (
            ("xl/workbook.xml", "x15 xr", '<extLst><ext uri="{synthetic}"><x15:future xr:uid="synthetic"/></ext></extLst>'),
            ("xl/worksheets/sheet4.xml", "x15 xr", '<extLst><ext uri="{synthetic}"><xr:revision/></ext></extLst>'),
        ):
            xml = package.read(name).decode("utf-8")
            xml = xml.replace("<workbook", f'<workbook{namespaces} mc:Ignorable="{ignorable}"', 1) if name.endswith("workbook.xml") else xml.replace("<worksheet", f'<worksheet{namespaces} mc:Ignorable="{ignorable}"', 1)
            xml = xml.replace("</workbook>" if name.endswith("workbook.xml") else "</worksheet>", extension + ("</workbook>" if name.endswith("workbook.xml") else "</worksheet>"))
            package.writestr(name, xml)


class V144ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.template = self.root / "synthetic-template.xlsx"; self.summary = self.root / "summary.xlsx"; self.output = self.root / "report.xlsx"
        make_template(self.template); make_summary(self.summary)

    def tearDown(self): self.temp.cleanup()

    def build(self, **kwargs):
        return build_v14_report(self.template, [self.summary], self.output, logging.getLogger("v14-test"), verify_hash=False, **kwargs)

    def test_mapping_clear_visibility_metadata_scalars_and_integrity(self):
        with zipfile.ZipFile(self.template) as before: protected = {name: before.read(name) for name in before.namelist() if name.startswith(("xl/charts/", "xl/drawings/", "xl/theme/", "xl/printerSettings/")) or name == "xl/styles.xml"}
        self.build()
        wb = load_workbook(self.output, data_only=False)
        self.assertEqual(wb.sheetnames, TARGETS + ["Metadata"])
        self.assertEqual(wb["Frequency Response_1_12"].sheet_state, "visible")
        self.assertEqual(wb["Frequency Response_orignal"].sheet_state, "visible")
        self.assertEqual(wb["Frequency Response_1_3"]["A40"].value, "SYNTH-0000")
        self.assertEqual(wb["Frequency Response_1_3"]["C40"].value, "PASS")
        self.assertEqual(wb["Frequency Response_1_12"]["C40"].value, "NOT_PROVIDED")
        self.assertIsNone(wb["Frequency Response_1_3"]["A41"].value)
        self.assertEqual(wb["SNR"]["B40"].value, 31.5); self.assertEqual(wb["Sensitivity"]["B40"].value, -31.5)
        self.assertEqual(wb["Metadata"]["C2"].value, "x0")
        with zipfile.ZipFile(self.output) as after:
            self.assertNotIn("xl/calcChain.xml", after.namelist())
            self.assertFalse(any(name.startswith("xl/externalLinks/") for name in after.namelist()))
            self.assertEqual(protected, {name: after.read(name) for name in protected})
            workbook = after.read("xl/workbook.xml").decode()
            self.assertIn('fullCalcOnLoad="1"', workbook); self.assertIn('forceFullCalc="1"', workbook)

    def test_compatibility_prefixes_and_existing_calcpr_survive(self):
        add_excel_compatibility_namespaces(self.template)
        self.build()
        with zipfile.ZipFile(self.output) as report:
            for name in ("xl/workbook.xml", "xl/worksheets/sheet4.xml"):
                xml = report.read(name).decode("utf-8")
                ignorable = __import__("re").search(r'Ignorable="([^"]+)"', xml).group(1).split()
                declared = set(__import__("re").findall(r'xmlns:([^=]+)=', xml))
                self.assertTrue(set(ignorable).issubset(declared), name)
            workbook = report.read("xl/workbook.xml").decode("utf-8")
            self.assertEqual(workbook.count("calcPr"), 1)

    def test_frequency_mismatch_is_fatal(self):
        make_summary(self.summary, mismatch=True)
        with self.assertRaisesRegex(V14ReportError, "Frequency-axis mismatch"): self.build()
        self.assertFalse(self.output.exists())

    def test_rawdata_not_provided_leaves_raw_frequency_cells_blank(self):
        make_summary(self.summary, rawdata_missing=True)
        self.build()
        for sheet in ("Frequency Response_1_12", "Frequency Response_orignal"):
            self.assertEqual(load_workbook(self.output)[sheet]["C40"].value, "NOT_PROVIDED")
            self.assertIsNone(load_workbook(self.output)[sheet]["D40"].value)

    def test_test_time_is_compact_template_text_and_clear_does_not_expand_grid(self):
        wb = load_workbook(self.summary)
        wb["03_FR_1_3"]["B2"] = datetime(2026, 9, 11, 12, 34, 56)
        wb.save(self.summary)
        self.build()
        self.assertEqual(load_workbook(self.output)["Frequency Response_1_3"]["B40"].value, "20260911123456")
        with zipfile.ZipFile(self.output) as report:
            self.assertNotIn(b'r="986"', report.read("xl/worksheets/sheet4.xml"))

    def test_capacity_and_chart_warning(self):
        make_summary(self.summary, rows=MAX_DUTS)
        with self.assertLogs("v14-test", "WARNING") as logs: self.build()
        self.assertTrue(any("Chart display coverage" in line for line in logs.output))
        self.assertEqual(load_workbook(self.output)["SNR"].cell(40 + MAX_DUTS - 1, 1).value, "SYNTH-0946")
        make_summary(self.summary, rows=MAX_DUTS + 1)
        with self.assertRaisesRegex(V14ReportError, "maximum"): self.build()

    def test_multiple_summaries_are_combined(self):
        second = self.root / "summary2.xlsx"; make_summary(second)
        build_v14_report(self.template, [self.summary, second], self.output, verify_hash=False)
        self.assertEqual(load_workbook(self.output)["Metadata"].max_row, 3)

    def test_production_entrypoint_runs_v13_then_fixed_template_report(self):
        archive = self.root / "ARUBA_MIC.zip"
        summary = self.root / "summary_MIC_Online.xlsx"
        with patch("src.v14_main.run_pipeline", return_value=[summary]) as pipeline, patch("src.v14_main.build_v14_report") as report:
            self.assertEqual(v14_main.main([str(archive), "--output-dir", str(self.root)]), 0)
        pipeline.assert_called_once()
        self.assertEqual(report.call_args.args[0], v14_main.DEFAULT_TEMPLATE)
        self.assertEqual(report.call_args.args[1], [summary])
        self.assertEqual(report.call_args.args[2], self.root / "report_MIC_Online.xlsx")

if __name__ == "__main__": unittest.main()
