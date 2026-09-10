import csv
import io
import logging
import subprocess
import tarfile
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from src.config import MatchConfig, SHEET_NAMES
from src.main import classify_inputs, run_pipeline


def csv_text(rows):
    stream = io.StringIO(newline="")
    csv.writer(stream, lineterminator="\r\n").writerows(rows)
    return stream.getvalue()


def main_rows(
    sn="APSYNTH001",
    run_id="000000000000000123",
    start="2026-09-10 16:30:00.125",
    end="2026-09-10 16:30:12.875",
    thd="1",
    phase="1",
    noise="1",
    snr="1",
    sensitivity="1",
    include_snr=True,
    include_sensitivity=True,
    station_id="ARUBA_MIC",
    tester="ARUBA-MIC-01",
    extra_items=(),
):
    rows = [
        ["uut_sn", "", sn, "", "", ""],
        ["tsr_id", "", run_id, "", "", ""],
        ["op_id", "", "OP-REDACTED", "", "", ""],
        ["station_id", "", station_id, "", "", ""],
        ["tester_id", "", tester, "", "", ""],
        ["test_sw_ver", "", "SYNTH-1.0", "", "", ""],
        ["test_start_time", "", start, "", "", ""],
        ["sfis_get_mac", "", "00:00:00:00:00:00", "", "", ""],
        ["FR100", "1", "-8.25", "", "", ""],
        ["THD", thd, "0.5", "", "", ""],
        ["Phase", phase, "2.0", "", "", ""],
        ["Noise", noise, "-70", "", "", ""],
    ]
    if include_snr:
        rows.append(["SNR", snr, "38.25", "", "", ""])
    if include_sensitivity:
        rows.append(["Sensitivity", sensitivity, "-31.5", "", "", ""])
    rows.extend([[key, "", value, "", "", ""] for key, value in extra_items])
    rows.extend([
        ["test_end_time", "", end, "", "", ""],
        ["total_test_time", "", "12.75", "", "", ""],
    ])
    return rows


def fr_rows(axis=(100, 200), marker=0):
    return [
        ["FR"],
        ["Test Result:", "FAIL"],
        ["Frequency(Hz):", *axis],
        ["Data (dBFS):", -10 + marker, -11 + marker],
        ["THD"],
        ["Test Result:", "PASS"],
        ["Frequency(Hz):", *axis],
        ["Data (%):", 0.1, 0.2],
        ["Phase"],
        ["Test Result:", "FAIL"],
        ["Frequency(Hz):", *axis],
        ["Data (Deg):", 1.0, 2.0],
    ]


def noise_rows(axis=(100, 200)):
    return [
        ["Noise"],
        ["Test Result:", "FAIL"],
        ["Frequency(Hz):", *axis],
        ["Data (dB(A)V):", -70, -71],
    ]


def raw_rows(axis=(100, 200)):
    return [
        ["FR_Original"],
        ["Frequency(Hz):", *axis],
        ["Data (dBFS):", -9, -10],
        ["FR_1/12smooth"],
        ["Frequency(Hz):", *axis],
        ["Data (dBFS):", -9.5, -10.5],
        ["FR_1/3smooth"],
        ["Frequency(Hz):", *axis],
        ["Data (dBFS):", 999, 999],
    ]


def add_run(
    files,
    run="run1",
    main_timestamp="202609101630001",
    fr_timestamp="20260910163010",
    noise_timestamp="20260910163020",
    sn="APSYNTH001",
    run_id="000000000000000123",
    path_result="PASS",
    fr_result="PASS",
    axis=(100, 200),
    test_type="MIC",
    mode="Online",
    **station_fields,
):
    station = f"ARUBA-{test_type}-01"
    base = f"{test_type}/{mode}/202609/20260910/{station}/{path_result}/{run}"
    station_fields.setdefault("station_id", f"ARUBA_{test_type}")
    station_fields.setdefault("tester", station)
    files[f"{base}/station_{main_timestamp}.csv"] = csv_text(
        main_rows(sn, run_id, **station_fields)
    )
    files[f"{base}/{sn}_FR_{fr_result}_{fr_timestamp}.csv"] = csv_text(
        fr_rows(axis)
    )
    files[f"{base}/{sn}_Noise_{noise_timestamp}.csv"] = csv_text(
        noise_rows(axis)
    )


def _write_staging(root, files):
    for name, content in files.items():
        target = root.joinpath(*name.split("/"))
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8")


def write_archive(path, files):
    lower = path.name.lower()
    if lower.endswith(".zip"):
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in files.items():
                archive.writestr(name, content)
        return
    if lower.endswith((".tar", ".tar.gz", ".tgz")):
        mode = "w:gz" if lower.endswith((".tar.gz", ".tgz")) else "w"
        with tarfile.open(path, mode) as archive:
            for name, content in files.items():
                payload = content.encode("utf-8")
                member = tarfile.TarInfo(name)
                member.size = len(payload)
                archive.addfile(member, io.BytesIO(payload))
        return
    if lower.endswith(".7z"):
        with tempfile.TemporaryDirectory() as directory:
            staging = Path(directory)
            _write_staging(staging, files)
            try:
                import py7zr
            except ImportError:
                subprocess.run(
                    ["tar", "--format", "7zip", "-cf", str(path), "-C", str(staging), "."],
                    check=True,
                    capture_output=True,
                )
            else:
                with py7zr.SevenZipFile(path, "w") as archive:
                    for name in files:
                        archive.write(staging.joinpath(*name.split("/")), arcname=name)
        return
    raise ValueError(path)


def headers(workbook, sheet):
    return [cell.value for cell in workbook[sheet][1]]


class V136PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.logger = logging.getLogger(f"test.{self.id()}")
        self.logger.addHandler(logging.NullHandler())

    def tearDown(self):
        self.temp.cleanup()

    def build(self, files, raw_files=None, suffix=".zip", config=None):
        mic = self.root / f"ARUBA_MIC{suffix}"
        write_archive(mic, files)
        raw = None
        if raw_files is not None:
            raw = self.root / f"RawData_RD{suffix}"
            write_archive(raw, raw_files)
        outputs = run_pipeline(
            mic, None, raw, self.root / "out", config, self.logger
        )
        self.assertEqual(len(outputs), 1)
        return load_workbook(outputs[0], data_only=False)

    def test_all_required_archive_formats(self):
        for suffix in (".zip", ".7z", ".tar", ".tar.gz", ".tgz"):
            with self.subTest(suffix=suffix):
                files = {}
                add_run(files)
                wb = self.build(files, suffix=suffix)
                self.assertEqual(tuple(wb.sheetnames), SHEET_NAMES)
                self.assertEqual(
                    classify_inputs([self.root / f"ARUBA_MIC{suffix}"])[0].suffixes,
                    (self.root / f"ARUBA_MIC{suffix}").suffixes,
                )

    def test_locked_contract_not_provided_and_scalar_results(self):
        files = {}
        add_run(
            files,
            path_result="FAIL",
            fr_result="PASS",
            thd="0",
            phase="1",
            noise="bad",
            snr="0",
            sensitivity="0",
        )
        wb = self.build(files)
        self.assertEqual(tuple(wb.sheetnames), SHEET_NAMES)
        self.assertEqual(
            headers(wb, "02_FR_original"),
            ["SN", "Test_Time", "RawData_Match_Status"],
        )
        self.assertEqual(
            headers(wb, "04_FR_1_12"),
            ["SN", "Test_Time", "RawData_Match_Status"],
        )
        self.assertEqual(wb["02_FR_original"]["C2"].value, "NOT_PROVIDED")
        self.assertIsNone(wb["02_FR_original"]["B2"].value)
        self.assertEqual(wb["05_THD"]["D2"].value, "FAIL")
        self.assertEqual(wb["06_Phase"]["D2"].value, "PASS")
        self.assertIsNone(wb["07_Noise"]["D2"].value)
        self.assertEqual(wb["08_SNR"]["D2"].value, "FAIL")
        self.assertEqual(wb["08_SNR"]["E2"].value, 38.25)
        self.assertEqual(wb["09_Sensitivity"]["D2"].value, "FAIL")
        self.assertEqual(wb["09_Sensitivity"]["E2"].value, -31.5)
        self.assertEqual(
            headers(wb, "08_SNR"),
            ["SN", "Test_Time", "Path_Result", "Result", "SNR_dB", "SNR_Source_Status"],
        )
        self.assertNotIn("Sensitivity_dBFS", headers(wb, "08_SNR"))
        workbook_headers = {
            value for sheet in wb.sheetnames for value in headers(wb, sheet)
        }
        self.assertNotIn("Section_Result", workbook_headers)
        self.assertNotIn("APx_Section_Result", workbook_headers)
        self.assertFalse({"N", "Mean", "Max", "Min", "Range", "STDEV"} & workbook_headers)
        self.assertEqual(wb["03_FR_1_3"]["C2"].value, "FAIL")
        self.assertEqual(wb["03_FR_1_3"]["D2"].value, "PASS")
        self.assertFalse(any(ws._charts for ws in wb.worksheets))

    def test_source_specific_times_and_fractional_station_times(self):
        files = {}
        add_run(files)
        raw = {
            "RawData_RD/MIC/APSYNTH001_FR_PASS_20260910163030.csv":
                csv_text(raw_rows())
        }
        wb = self.build(files, raw)
        expected = {
            "01_Metadata": datetime(2026, 9, 10, 16, 30, 0),
            "02_FR_original": datetime(2026, 9, 10, 16, 30, 30),
            "03_FR_1_3": datetime(2026, 9, 10, 16, 30, 10),
            "04_FR_1_12": datetime(2026, 9, 10, 16, 30, 30),
            "05_THD": datetime(2026, 9, 10, 16, 30, 10),
            "06_Phase": datetime(2026, 9, 10, 16, 30, 10),
            "07_Noise": datetime(2026, 9, 10, 16, 30, 20),
            "08_SNR": datetime(2026, 9, 10, 16, 30, 0),
            "09_Sensitivity": datetime(2026, 9, 10, 16, 30, 0),
        }
        metadata_headers = headers(wb, "01_Metadata")
        metadata_row = dict(
            zip(metadata_headers, [cell.value for cell in wb["01_Metadata"][2]])
        )
        self.assertEqual(metadata_row["Test_Time"], expected["01_Metadata"])
        self.assertEqual(
            metadata_row["test_start_time"],
            datetime(2026, 9, 10, 16, 30, 0, 125000),
        )
        self.assertEqual(
            metadata_row["test_end_time"],
            datetime(2026, 9, 10, 16, 30, 12, 875000),
        )
        for sheet, timestamp in expected.items():
            if sheet != "01_Metadata":
                self.assertEqual(wb[sheet]["B2"].value, timestamp, sheet)
        self.assertEqual(
            headers(wb, "02_FR_original"),
            ["SN", "Test_Time", "RawData_Match_Status", 100, 200],
        )
        self.assertEqual(
            headers(wb, "04_FR_1_12"),
            ["SN", "Test_Time", "RawData_Match_Status", 100, 200],
        )

    def test_metadata_complete_capture_union_and_identifier_text(self):
        files = {}
        add_run(
            files,
            run="first",
            run_id="000000000000000123",
            extra_items=(("custom_alpha", "0000456"),),
        )
        add_run(
            files,
            run="second",
            main_timestamp="20260910163100",
            fr_timestamp="20260910163110",
            noise_timestamp="20260910163120",
            run_id="000000000000000124",
            extra_items=(("custom_beta", "beta-value"),),
        )
        wb = self.build(files)
        sheet = wb["01_Metadata"]
        header = headers(wb, "01_Metadata")
        self.assertTrue({
            "uut_sn", "tsr_id", "op_id", "station_id", "tester_id",
            "test_sw_ver", "test_start_time", "sfis_get_mac", "FR100",
            "THD", "Phase", "Noise", "SNR", "Sensitivity",
            "test_end_time", "total_test_time", "custom_alpha", "custom_beta",
        }.issubset(header))
        self.assertLess(header.index("custom_alpha"), header.index("custom_beta"))
        rows = [
            dict(zip(header, [cell.value for cell in row]))
            for row in sheet.iter_rows(min_row=2)
        ]
        first = next(row for row in rows if row["tsr_id"] == "000000000000000123")
        second = next(row for row in rows if row["tsr_id"] == "000000000000000124")
        self.assertEqual(first["custom_alpha"], "0000456")
        self.assertIsNone(first["custom_beta"])
        self.assertIsNone(second["custom_alpha"])
        self.assertEqual(second["custom_beta"], "beta-value")
        self.assertEqual(first["sfis_get_mac"], "00:00:00:00:00:00")
        self.assertIsInstance(first["FR100"], float)

    def test_rawdata_unmatched_and_ambiguous(self):
        files = {}
        add_run(files)
        far = {
            "RawData_RD/MIC/APSYNTH001_FR_PASS_20260910163300.csv":
                csv_text(raw_rows())
        }
        wb = self.build(files, far)
        metadata = dict(
            zip(
                headers(wb, "01_Metadata"),
                [cell.value for cell in wb["01_Metadata"][2]],
            )
        )
        self.assertEqual(metadata["RawData_Match_Status"], "UNMATCHED")
        near = {
            "RawData_RD/MIC/APSYNTH001_FR_PASS_20260910163008.csv":
                csv_text(raw_rows()),
            "RawData_RD/MIC/APSYNTH001_FR_PASS_20260910163012.csv":
                csv_text(raw_rows()),
        }
        wb = self.build(files, near, config=MatchConfig(60, 5))
        self.assertEqual(wb["02_FR_original"]["C2"].value, "AMBIGUOUS")
        self.assertIsNone(wb["02_FR_original"]["D2"].value)

    def test_latest_duplicate_selection_and_axis_mismatch(self):
        files = {}
        add_run(files, run="old", run_id="RUN-1")
        add_run(
            files,
            run="new",
            main_timestamp="202609101631001",
            fr_timestamp="20260910163110",
            noise_timestamp="20260910163120",
            run_id="RUN-2",
            axis=(100, 300),
        )
        base = "MIC/Online/202609/20260910/ARUBA-MIC-01/PASS/old"
        files[f"{base}/APSYNTH001_FR_PASS_20260910163040.csv"] = csv_text(
            fr_rows(marker=5)
        )
        wb = self.build(files)
        header = headers(wb, "01_Metadata")
        rows = [
            dict(zip(header, [cell.value for cell in row]))
            for row in wb["01_Metadata"].iter_rows(min_row=2)
        ]
        self.assertEqual(sum(row["Latest_Run"] == "TRUE" for row in rows), 1)
        self.assertEqual(
            next(row for row in rows if row["Latest_Run"] == "TRUE")["Run_ID"],
            "RUN-2",
        )
        self.assertIn(
            "20260910163040",
            next(row for row in rows if row["Run_ID"] == "RUN-1")["FR_CSV"],
        )
        curve_rows = [
            [cell.value for cell in row[4:]]
            for row in wb["03_FR_1_3"].iter_rows(min_row=2)
        ]
        self.assertEqual(
            sum(all(value is None for value in values) for values in curve_rows), 1
        )

    def test_missing_snr_and_sensitivity_are_blank(self):
        files = {}
        add_run(files, include_snr=False, include_sensitivity=False)
        wb = self.build(files)
        for sheet in ("08_SNR", "09_Sensitivity"):
            row = [cell.value for cell in wb[sheet][2]]
            self.assertIsNone(row[3])
            self.assertIsNone(row[4])
            self.assertEqual(row[5], "Missing in Main Station CSV")

    def test_invalid_sensitivity_result_is_blank_without_substitute(self):
        files = {}
        add_run(files, sensitivity="invalid")
        wb = self.build(files)
        self.assertIsNone(wb["09_Sensitivity"]["D2"].value)
        self.assertEqual(wb["09_Sensitivity"]["E2"].value, -31.5)
        self.assertEqual(
            wb["09_Sensitivity"]["F2"].value,
            "Present in Main Station CSV",
        )

    def test_shared_mic_premic_pipeline_and_offline_discovery(self):
        mic_files = {}
        premic_files = {}
        add_run(mic_files)
        add_run(
            premic_files,
            test_type="PREMIC",
            mode="Offline",
            sn="APSYNTH002",
            run_id="RUN-SYNTH-PREMIC",
        )
        mic = self.root / "ARUBA_MIC.zip"
        premic = self.root / "ARUBA_PREMIC.tgz"
        write_archive(mic, mic_files)
        write_archive(premic, premic_files)
        outputs = run_pipeline(
            mic, premic, None, self.root / "out", logger=self.logger
        )
        self.assertEqual(
            {path.name for path in outputs},
            {"summary_MIC_Online.xlsx", "summary_PREMIC_Offline.xlsx"},
        )

    def test_requires_at_least_one_fatp_archive(self):
        with self.assertRaisesRegex(ValueError, "At least one FATP archive"):
            run_pipeline(None, None, None, self.root, logger=self.logger)


if __name__ == "__main__":
    unittest.main()

